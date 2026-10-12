"""Safe HTML and Excel exports for two saved execution histories."""
from collections import Counter
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .run_comparison import compare_runs

HEADERS = ["API", "Method", "순번", "기준 결과", "비교 결과", "변화", "기준 HTTP", "비교 HTTP", "시간차(ms)"]
FIELDS = ["name", "method", "occurrence", "before", "after", "change", "before_http", "after_http", "delta_ms"]
CHANGES = ("RECOVERED", "REGRESSED", "CHANGED", "UNCHANGED", "ADDED", "REMOVED")
FILLS = {"RECOVERED": "D9EAD3", "REGRESSED": "F4CCCC", "CHANGED": "FFF2CC",
         "ADDED": "D9EAF7", "REMOVED": "E7E6E6"}


def safe_excel_value(value):
    # User-supplied Postman names and URLs must never become spreadsheet formulas.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def summarize(changes):
    counts = Counter(row["change"] for row in changes)
    return {kind: counts[kind] for kind in CHANGES}


def export_comparison(storage, before_run_id: str, after_run_id: str,
                      directory: Path) -> list[Path]:
    if before_run_id == after_run_id:
        raise ValueError("서로 다른 회차를 선택하세요.")
    directory = Path(directory)
    changes = compare_runs(storage.results(before_run_id), storage.results(after_run_id))
    counts = summarize(changes)
    directory.mkdir(parents=True, exist_ok=True)
    prefix = f"compare_{before_run_id}_{after_run_id}"
    xlsx_path = directory / f"{prefix}.xlsx"
    html_path = directory / f"{prefix}.html"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "변화 상세"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in changes:
        sheet.append([safe_excel_value(row[field]) for field in FIELDS])
        color = FILLS.get(row["change"])
        if color:
            sheet.cell(sheet.max_row, 6).fill = PatternFill("solid", fgColor=color)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 36, "B": 12, "C": 10, "D": 15, "E": 15,
                          "F": 19, "G": 15, "H": 15, "I": 19}.items():
        sheet.column_dimensions[column].width = width
    summary_sheet = workbook.create_sheet("변화 집계")
    summary_sheet.append(["분류", "건수"])
    for kind, count in counts.items():
        summary_sheet.append([kind, count])
    workbook.save(xlsx_path)

    summary_rows = "".join(f"<tr><td>{kind}</td><td>{count}</td></tr>"
                           for kind, count in counts.items())
    details = "".join(
        "<tr>" + "".join(f"<td>{escape(str(row[field] if row[field] is not None else ''))}</td>"
                          for field in FIELDS) + "</tr>" for row in changes)
    html_path.write_text(
        "<!doctype html><html lang='ko'><meta charset='utf-8'>"
        "<title>API 회차 비교</title><style>"
        "body{font:14px sans-serif;margin:32px}table{border-collapse:collapse;width:100%}"
        "th,td{border:1px solid #ddd;padding:8px}th{background:#edf1f5}</style>"
        "<h1>API 실행 회차 비교</h1>"
        f"<p>기준 회차: {escape(before_run_id)} / 비교 회차: {escape(after_run_id)}</p>"
        "<h2>변화 집계</h2><table><tr><th>변화</th><th>건수</th></tr>"
        + summary_rows + "</table><h2>API별 변화</h2><table><tr>"
        + "".join(f"<th>{escape(title)}</th>" for title in HEADERS) + "</tr>"
        + details + "</table></html>", encoding="utf-8")
    return [xlsx_path, html_path]
