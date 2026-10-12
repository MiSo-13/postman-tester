"""Regression coverage for cancellation boundaries and exported report text."""
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from app.main_window import RunWorker
from app.reports import export_reports, safe_spreadsheet_value
from app.retry_policy import RetryPolicy


class StorageStub:
    def __init__(self):
        self.rows = []
        self.final = None
        self.saved = []

    def results(self, run_id):
        return [{
            "endpoint_name": "=1+1",
            "method": "GET",
            "url": "https://example.test",
            "status": "FAIL",
            "http_status": 500,
            "elapsed_ms": 1.0,
            "detail": "@SUM(A1:A3)\x01"
        }]

    def save_report(self, run_id, fmt, path):
        self.saved.append((fmt, path))

    def add_result(self, run_id, name, method, url, **result):
        self.rows.append(result)

    def finish_run(self, run_id, status):
        self.final = status


def test_xlsx_imported_strings_do_not_execute_as_formulas(tmp_path):
    storage = StorageStub()
    xlsx_path, html_path = export_reports(storage, "report-1", tmp_path)
    wb = load_workbook(xlsx_path)
    ws = wb.active
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value == "'=1+1"
    assert ws["G2"].data_type == "s"
    assert ws["G2"].value == "'@SUM(A1:A3)"
    assert "&" not in html_path.read_text(encoding="utf-8").split("report-1")[0]
    assert safe_spreadsheet_value(42) == 42


def test_cancel_before_login_does_not_contact_server():
    storage = StorageStub()
    auth_endpoint = SimpleNamespace(
        method="GET", name="profile", url="{{baseUrl}}/me",
        request={"auth": {"type": "bearer", "bearer": [
            {"key": "token", "value": "{{userAccessToken}}"}]}})
    settings = SimpleNamespace(
        values={"default_timeout_seconds": 1, "verify_ssl": True,
                "dangerous_methods_enabled": False, "file_root": "data/uploads"},
        root=Path("."))
    worker = RunWorker(storage, "run-id", [("a" * 64, auth_endpoint)],
                       {"baseUrl": "https://example.test"}, settings, {},
                       user_credentials=("id", "pw"), retry_policy=RetryPolicy(0, 0))
    worker.cancel()
    with patch("app.main_window.login_tokens") as login, \
         patch("app.main_window.execute") as execute:
        worker.run()
    login.assert_not_called()
    execute.assert_not_called()
    assert storage.final == "CANCELLED"
    assert storage.rows == []


def test_cancel_during_login_prevents_followup_requests():
    storage = StorageStub()
    auth_endpoint = SimpleNamespace(
        method="GET", name="profile", url="{{baseUrl}}/me",
        request={"auth": {"type": "bearer", "bearer": [
            {"key": "token", "value": "{{userAccessToken}}"}]}})
    settings = SimpleNamespace(
        values={"default_timeout_seconds": 1, "verify_ssl": True,
                "dangerous_methods_enabled": False, "file_root": "data/uploads"},
        root=Path("."))
    worker = RunWorker(storage, "run-id", [("a" * 64, auth_endpoint)],
                       {"baseUrl": "https://example.test"}, settings, {},
                       user_credentials=("id", "pw"), retry_policy=RetryPolicy(0, 0))
    def login(*args, **kwargs):
        worker.cancel()
        return "access", "refresh"
    with patch("app.main_window.login_tokens", side_effect=login), \
         patch("app.main_window.execute") as execute:
        worker.run()
    execute.assert_not_called()
    assert storage.final == "CANCELLED"
