"""Durable HTML/XLSX exports for saved test runs."""
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


COLUMNS = ["endpoint_name", "method", "url", "status", "http_status", "elapsed_ms", "detail"]
HEADERS = ["API 이름", "Method", "URL", "결과", "HTTP 상태", "응답(ms)", "상세"]


def safe_spreadsheet_value(value):
    """Treat imported API text and response details as data, never spreadsheet formulas."""
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_reports(storage, run_id: str, output_directory: Path) -> list[Path]:
    output_directory.mkdir(parents=True, exist_ok=True)
    rows = storage.results(run_id)
    xlsx_path = output_directory / f"{run_id}.xlsx"
    html_path = output_directory / f"{run_id}.html"
    wb = Workbook()
    ws = wb.active
    ws.title = "API Results"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([safe_spreadsheet_value(row[key]) for key in COLUMNS])
        color = {"PASS": "D9EAD3", "FAIL": "F4CCCC", "ERROR": "FCE5CD",
                 "BLOCKED": "FFF2CC", "SKIPPED": "E7E6E6"}.get(row["status"])
        if color:
            ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=color)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, width in {"A": 35, "B": 12, "C": 55, "D": 16, "E": 16, "F": 18, "G": 75}.items():
        ws.column_dimensions[column].width = width
    wb.save(xlsx_path)
    html_rows = "".join("<tr>" + "".join(f"<td>{escape(str(row[key] if row[key] is not None else ''))}</td>"
                        for key in COLUMNS) + "</tr>" for row in rows)
    run = storage.get_run(run_id)
    run_status = run["status"] if run else "UNKNOWN"
    counts = {status: sum(row["status"] == status for row in rows)
              for status in ("PASS", "FAIL", "ERROR", "BLOCKED", "SKIPPED")}
    empty_note = (
        "<p class='warning'>API 실행 결과가 없습니다. 로그인 실패, 실행 취소 또는 "
        "실행 전 오류가 발생했을 수 있습니다. 앱의 실행 로그를 확인하세요.</p>"
        if not rows else ""
    )
    html_path.write_text(
        "<!doctype html><html lang='ko'><meta charset='utf-8'>"
        "<title>API 테스트 결과</title><style>body{font:14px sans-serif;margin:32px;background:#f6f8fc;color:#202b3a;}"
        "table{border-collapse:collapse;width:100%}td,th{border:1px solid #ddd;padding:8px}"
        "th{background:#edf1f5}tr:hover{background:#f5f8ff}"
        ".summary{background:white;padding:16px;border:1px solid #dce4ed;border-radius:8px;margin-bottom:16px}"
        ".warning{padding:14px;background:#fff1dc;border-left:4px solid #df8d19}</style>"
        f"<h1>API 테스트 결과</h1><div class='summary'>Run ID: {escape(run_id)}"
        f"<p>실행 상태: <b>{escape(run_status)}</b> · 요청 결과: {len(rows)}건 · "
        + " · ".join(f"{key} {value}" for key, value in counts.items()) + "</p></div>"
        + empty_note +
        "<table><thead><tr>" + "".join(f"<th>{escape(h)}</th>" for h in HEADERS)
        + "</tr></thead><tbody>" + html_rows + "</tbody></table></html>",
        encoding="utf-8")
    for fmt, path in (("xlsx", xlsx_path), ("html", html_path)):
        storage.save_report(run_id, fmt, path)
    return [xlsx_path, html_path]
