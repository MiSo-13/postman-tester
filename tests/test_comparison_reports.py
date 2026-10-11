"""Regression tests for comparison reports: totals, escaping, and Excel safety."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.comparison_reports import export_comparison, summarize


def saved(name, status, ms=20, http=200):
    return {"endpoint_name": name, "method": "GET",
            "url": "https://example.test/resource",
            "status": status, "http_status": http,
            "elapsed_ms": ms, "detail": "SECRET_DO_NOT_EXPORT"}


class FakeStorage:
    def __init__(self, left, right):
        self.data = {"before": left, "after": right}

    def results(self, run_id):
        return self.data[run_id]


def test_export_comparison_html_xlsx_and_counts(tmp_path):
    storage = FakeStorage(
        [saved("user", "FAIL", 50), saved("old", "PASS", 30)],
        [saved("user", "PASS", 35), saved("new", "FAIL", 60)])
    paths = export_comparison(storage, "before", "after", tmp_path / "reports")
    assert len(paths) == 2
    assert all(path.exists() for path in paths)
    workbook = load_workbook(paths[0], data_only=True)
    assert workbook.sheetnames == ["변화 상세", "변화 집계"]
    counts = {row[0]: row[1] for row in
              list(workbook["변화 집계"].values)[1:]}
    assert counts["RECOVERED"] == 1
    assert counts["ADDED"] == 1
    assert counts["REMOVED"] == 1
    assert counts["REGRESSED"] == 0
    text = paths[1].read_text(encoding="utf-8")
    assert "RECOVERED" in text
    assert "SECRET_DO_NOT_EXPORT" not in text


def test_escapes_html_and_blocks_excel_formula_injection(tmp_path):
    name = "=2+2<script>alert(1)</script>"
    storage = FakeStorage([saved(name, "FAIL")], [saved(name, "PASS")])
    paths = export_comparison(storage, "before", "after", tmp_path)
    book = load_workbook(paths[0])
    cell = book["변화 상세"]["A2"]
    assert cell.data_type != "f"
    assert cell.value.startswith("'=")
    markup = paths[1].read_text(encoding="utf-8")
    assert "&lt;script&gt;" in markup
    assert "<script>" not in markup


def test_same_run_is_rejected(tmp_path):
    with pytest.raises(ValueError):
        export_comparison(FakeStorage([], []), "before", "before", tmp_path)
    assert not (tmp_path / "compare_before_before.xlsx").exists()


def test_summary_has_all_classes_even_when_empty():
    assert summarize([]) == {
        "RECOVERED": 0, "REGRESSED": 0, "CHANGED": 0,
        "UNCHANGED": 0, "ADDED": 0, "REMOVED": 0}
